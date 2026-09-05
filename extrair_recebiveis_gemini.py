#!/usr/bin/env python3
"""Extrai tabelas visiveis de um PDF de recebiveis com Gemini e gera XLSX.

Uso normal:
    python extrair_recebiveis_gemini.py relatorio.pdf -o relatorio_tabelas.xlsx

Uso offline, a partir de um JSON previamente extraido:
    python extrair_recebiveis_gemini.py --json extracao.json -o relatorio_tabelas.xlsx

A chave deve existir somente na variavel de ambiente GEMINI_API_KEY.
"""

from __future__ import annotations

import argparse
import base64
import json
import math
import os
import re
import sys
import time
import unicodedata
from datetime import date, datetime
from decimal import Decimal, InvalidOperation
from pathlib import Path
from typing import Literal

from openpyxl import Workbook, load_workbook
from openpyxl.cell import Cell
from openpyxl.styles import Alignment, Border, Font, PatternFill, Side
from openpyxl.utils import get_column_letter
from openpyxl.worksheet.table import Table, TableStyleInfo
from pydantic import BaseModel, ConfigDict, Field, ValidationError, model_validator
from pypdf import PdfReader


APP_VERSION = "1.0.0"
DEFAULT_MODEL = os.getenv("GEMINI_MODEL", "gemini-3.8-flash")


class UserFacingError(RuntimeError):
    """Mensagem controlada, segura para apresentar ao colaborador."""

# Padrao visual reproduzido dos arquivos entregues anteriormente.
FUNDACRED_BLUE = "FF00308F"
SOURCE_BLUE = "FFE8EEF8"
HEADER_GRAY = "FFD9DEE7"
TOTAL_GRAY = "FFD1D5DB"
TEXT_DARK = "FF1F2937"
TEXT_MUTED = "FF5F6B7A"
WHITE = "FFFFFFFF"
BORDER_GRAY = "FFAAB4C3"

DataType = Literal[
    "text",
    "integer",
    "decimal",
    "currency",
    "percentage",
    "date",
    "datetime",
    "boolean",
]
RowKind = Literal["data", "subtotal", "total", "note"]


class ColumnDefinition(BaseModel):
    model_config = ConfigDict(extra="forbid")

    name: str = Field(description="Nome visivel da coluna no PDF.")
    data_type: DataType = Field(
        description="Tipo semantico usado para converter o valor no Excel."
    )
    number_format: str = Field(
        default="",
        description="Formato sugerido; pode ficar vazio porque o backend aplica um padrao.",
    )


class TableRow(BaseModel):
    model_config = ConfigDict(extra="forbid")

    row_kind: RowKind = Field(description="Classificacao da linha.")
    values: list[str] = Field(
        description=(
            "Valores exatamente como aparecem no PDF. Use string vazia para celula vazia."
        )
    )


class ExtractedTable(BaseModel):
    model_config = ConfigDict(extra="forbid")

    table_id: str
    sheet_name: str
    title: str
    source_pages: list[int]
    source_note: str
    limitations: list[str]
    columns: list[ColumnDefinition]
    rows: list[TableRow]

    @model_validator(mode="after")
    def validate_shape(self) -> "ExtractedTable":
        if not self.columns:
            raise ValueError(f"A tabela {self.table_id!r} nao possui colunas.")

        column_count = len(self.columns)
        for index, row in enumerate(self.rows, start=1):
            if len(row.values) != column_count:
                raise ValueError(
                    f"Tabela {self.table_id!r}, linha {index}: "
                    f"esperadas {column_count} celulas, recebidas {len(row.values)}."
                )
        return self


class DocumentExtraction(BaseModel):
    model_config = ConfigDict(extra="forbid")

    source_file: str
    page_count: int
    warnings: list[str]
    tables: list[ExtractedTable]


EXTRACTION_PROMPT = """
Voce esta extraindo dados de um relatorio financeiro de recebiveis exportado em PDF,
normalmente originado no Power BI. Sua resposta sera validada por um programa e usada
para criar um arquivo Excel. Analise o documento inteiro e devolva exclusivamente o
JSON solicitado pelo schema.

OBJETIVO
- Identifique TODAS as tabelas e blocos de indicadores visiveis no PDF.
- Cada tabela logica deve resultar em uma entrada separada em `tables`.
- Preserve a ordem em que as tabelas aparecem no documento.
- Blocos de KPIs relacionados podem formar uma tabela de duas colunas, por exemplo
  `Indicador` e `Valor`.

TABELAS QUE COSTUMAM EXISTIR
- Resumo de processos ativos, prorrogados, reparcelados e antecipados.
- Inadimplencia por janela de tempo.
- Recuperacao de valores e parcelas por ano/mes.
- Recebiveis futuros.
- Repasse mensal.
- Agregado mensal.
- Valores vencidos.
- Posicao de uma data de referencia.
- Valores a vencer e vencidos por ano.
Essa lista e apenas uma referencia. Nao invente uma tabela ausente e nao ignore uma
tabela diferente que esteja realmente visivel.

REGRAS DE TRANSCRICAO
- Copie textos, rotulos, datas, sinais, percentuais e valores conforme aparecem.
- Em `rows.values`, devolva strings com o valor visivel. Exemplo: `R$ 1.234,56`,
  `4,34%`, `31/08/2026`, `2026`. O Python convertera os tipos depois.
- Use string vazia para celulas vazias; nunca use marcadores inventados.
- Identificadores, contratos, documentos e codigos com zeros a esquerda devem ser
  classificados como `text`.
- Classifique colunas financeiras como `currency`, percentuais como `percentage`,
  quantidades inteiras como `integer`, medidas numericas como `decimal`, datas como
  `date` e periodos textuais como `text`.
- Linhas comuns usam `row_kind=data`; subtotais usam `subtotal`; totais exibidos no
  documento usam `total`; observacoes que fazem parte da grade usam `note`.
- Preserve os totais oficiais mostrados no PDF mesmo quando diferirem alguns centavos
  da soma das linhas. Registre a divergencia em `limitations` ou `warnings`.

CONTINUACAO E DUPLICIDADE
- Una tabelas que continuam em paginas seguintes quando titulo e colunas forem os mesmos.
- Nao repita cabecalhos que aparecem novamente por causa da paginacao.
- Nao duplique a mesma tabela quando existir uma visualizacao resumida e outra detalhada;
  prefira a versao mais completa e registre a decisao em `limitations`.

LIMITACOES VISUAIS
- Extraia somente o conteudo efetivamente visivel/renderizado no PDF.
- Uma barra de rolagem indica que podem existir linhas ocultas. Nao invente nem estime
  essas linhas; extraia apenas as visiveis e registre essa limitacao.
- Nao transforme graficos, decoracoes ou textos soltos em tabelas, salvo quando forem
  claramente KPIs tabulares.
- Se uma celula estiver ilegivel, deixe-a vazia e registre a ocorrencia.

METADADOS
- `source_pages` usa numeracao iniciando em 1.
- `sheet_name` deve ser curto, humano e representativo; o Python corrigira caracteres
  invalidos, duplicidades e o limite de 31 caracteres.
- `source_note` deve ser curta e descrever a origem ou o escopo da tabela.
- `page_count` deve refletir o total de paginas do PDF.
- Nao inclua explicacoes fora do JSON.
""".strip()


def parse_args() -> argparse.Namespace:
    parser = argparse.ArgumentParser(
        description="Extrai todas as tabelas visiveis de um PDF e gera XLSX."
    )
    source = parser.add_mutually_exclusive_group(required=True)
    source.add_argument("pdf", nargs="?", type=Path, help="PDF de entrada.")
    source.add_argument(
        "--json",
        dest="json_input",
        type=Path,
        help="JSON estruturado para gerar o Excel sem chamar o Gemini.",
    )
    parser.add_argument("-o", "--output", type=Path, help="Arquivo XLSX de saida.")
    parser.add_argument(
        "--model",
        default=DEFAULT_MODEL,
        help=f"Modelo Gemini. Padrao: {DEFAULT_MODEL}",
    )
    parser.add_argument(
        "--save-json",
        type=Path,
        help="Salva o JSON validado para auditoria.",
    )
    parser.add_argument("--max-pages", type=int, default=10)
    parser.add_argument("--max-mb", type=float, default=10.0)
    parser.add_argument("--attempts", type=int, default=2)
    parser.add_argument(
        "--no-freeze",
        action="store_true",
        help="Nao congela as quatro primeiras linhas das abas.",
    )
    parser.add_argument("--version", action="version", version=APP_VERSION)
    return parser.parse_args()


def default_output_path(pdf_path: Path | None, json_path: Path | None) -> Path:
    source = pdf_path or json_path
    assert source is not None
    return source.with_name(f"{source.stem}_tabelas.xlsx")


def validate_pdf(pdf_path: Path, max_pages: int, max_mb: float) -> int:
    if not pdf_path.exists() or not pdf_path.is_file():
        raise FileNotFoundError(f"PDF nao encontrado: {pdf_path}")
    if pdf_path.suffix.lower() != ".pdf":
        raise ValueError("O arquivo de entrada deve possuir extensao .pdf.")

    size_mb = pdf_path.stat().st_size / (1024 * 1024)
    if size_mb > max_mb:
        raise ValueError(
            f"PDF possui {size_mb:.2f} MB; o limite configurado e {max_mb:.2f} MB."
        )

    with pdf_path.open("rb") as stream:
        if stream.read(5) != b"%PDF-":
            raise ValueError("O conteúdo não possui assinatura PDF válida.")
    try:
        reader = PdfReader(str(pdf_path))
    except Exception:
        raise UserFacingError("PDF inválido ou danificado. Exporte novamente o relatório.") from None
    if reader.is_encrypted:
        raise ValueError("PDF protegido por senha nao e aceito.")
    page_count = len(reader.pages)
    if page_count == 0:
        raise ValueError("O PDF nao possui paginas.")
    if page_count > max_pages:
        raise ValueError(
            f"PDF possui {page_count} paginas; o limite configurado e {max_pages}."
        )
    return page_count


def extract_with_gemini(
    pdf_path: Path,
    model: str,
    expected_pages: int,
    attempts: int,
    api_key: str | None = None,
) -> DocumentExtraction:
    api_key = api_key or os.getenv("GEMINI_API_KEY")
    if not api_key:
        raise RuntimeError(
            "Defina GEMINI_API_KEY no ambiente. Nao coloque a chave no codigo."
        )

    try:
        from google import genai
    except ImportError as exc:
        raise RuntimeError(
            "Pacote google-genai ausente. Execute: pip install -r requirements.txt"
        ) from exc

    client = genai.Client(api_key=api_key)
    pdf_b64 = base64.b64encode(pdf_path.read_bytes()).decode("ascii")
    prompt = (
        f"Arquivo: {pdf_path.name}\n"
        f"Quantidade de paginas validada pelo backend: {expected_pages}\n\n"
        f"{EXTRACTION_PROMPT}"
    )

    last_error: Exception | None = None
    for attempt in range(1, max(1, attempts) + 1):
        try:
            interaction = client.interactions.create(
                model=model,
                store=False,
                input=[
                    {
                        "type": "document",
                        "data": pdf_b64,
                        "mime_type": "application/pdf",
                    },
                    {"type": "text", "text": prompt},
                ],
                response_format={
                    "type": "text",
                    "mime_type": "application/json",
                    "schema": DocumentExtraction.model_json_schema(),
                },
            )
            output_text = getattr(interaction, "output_text", None)
            if not output_text:
                raise RuntimeError("O Gemini nao retornou JSON na resposta.")

            extraction = DocumentExtraction.model_validate_json(output_text)
            if extraction.page_count != expected_pages:
                extraction.warnings.append(
                    "O Gemini informou quantidade de paginas diferente do PDF; "
                    "foi mantida a contagem validada pelo backend."
                )
                extraction.page_count = expected_pages
            extraction.source_file = pdf_path.name
            client.close()
            return extraction
        except Exception as exc:  # Erros HTTP/SDK variam entre versoes.
            last_error = exc
            status = getattr(exc, "status_code", getattr(exc, "code", None))
            if status in (400, 401, 403, 404):
                break
            if attempt < attempts:
                time.sleep(min(2**attempt, 8))

    client.close()
    assert last_error is not None
    messages = {
        400: "A API rejeitou a solicitação. Verifique GEMINI_MODEL e a compatibilidade com PDF e JSON.",
        401: "Credencial Gemini inválida. Solicite ao administrador a revisão dos Secrets.",
        403: "Acesso à Gemini negado. Solicite ao administrador a revisão das permissões.",
        404: "Modelo indisponível nesta conta. Configure GEMINI_MODEL nos Secrets ou no ambiente, sem editar o código.",
        429: "Limite da Gemini atingido. Aguarde e tente novamente; o administrador deve verificar a cota.",
    }
    raise UserFacingError(messages.get(status, "Não foi possível obter uma extração válida da Gemini. Tente novamente.")) from None


def load_extraction_json(json_path: Path) -> DocumentExtraction:
    if not json_path.exists():
        raise FileNotFoundError(f"JSON nao encontrado: {json_path}")
    return DocumentExtraction.model_validate_json(json_path.read_text(encoding="utf-8"))


def strip_accents(value: str) -> str:
    normalized = unicodedata.normalize("NFKD", value)
    return "".join(char for char in normalized if not unicodedata.combining(char))


def unique_sheet_name(name: str, used: set[str]) -> str:
    cleaned = re.sub(r"[\\/*?:\[\]]", "-", name or "Tabela").strip(" .'")
    cleaned = re.sub(r"\s+", " ", cleaned) or "Tabela"
    cleaned = cleaned[:31]
    candidate = cleaned
    suffix = 2
    while candidate.casefold() in used:
        marker = f" ({suffix})"
        candidate = f"{cleaned[: 31 - len(marker)]}{marker}"
        suffix += 1
    used.add(candidate.casefold())
    return candidate


def unique_table_name(sheet_name: str, index: int, used: set[str]) -> str:
    base = re.sub(r"[^A-Za-z0-9_]", "", strip_accents(sheet_name).replace(" ", "_"))
    if not base or not base[0].isalpha():
        base = f"Tabela_{base}"
    base = f"Tbl_{base}"[:200]
    candidate = base
    suffix = 2
    while candidate.casefold() in used:
        candidate = f"{base[:190]}_{suffix}"
        suffix += 1
    used.add(candidate.casefold())
    return candidate


def clean_visible_number(value: str) -> tuple[str, bool]:
    text = value.strip().replace("\u00a0", " ")
    if not text or text in {"-", "--", "n/a", "N/A", "null", "NULL"}:
        return "", False

    negative = text.startswith("(") and text.endswith(")")
    text = text.strip("()")
    text = re.sub(r"(?i)R\$|US\$|USD|BRL", "", text)
    text = text.replace("%", "").replace(" ", "")
    if not re.fullmatch(r"[+-]?\d[\d,.]*", text):
        raise ValueError("Valor numérico inválido.")
    return text, negative


def parse_decimal(value: str, integer: bool = False) -> Decimal | None:
    text, parenthesis_negative = clean_visible_number(value)
    if not text:
        return None

    if integer:
        normalized = re.sub(r"[.,](?=\d{3}(?:[.,]|$))", "", text)
        normalized = normalized.replace(",", ".")
    elif "," in text and "." in text:
        if text.rfind(",") > text.rfind("."):
            normalized = text.replace(".", "").replace(",", ".")
        else:
            normalized = text.replace(",", "")
    elif "," in text:
        normalized = text.replace(".", "").replace(",", ".")
    else:
        normalized = text.replace(",", "")

    try:
        number = Decimal(normalized)
    except InvalidOperation as exc:
        raise ValueError(f"Valor numerico invalido: {value!r}") from exc
    if parenthesis_negative and number > 0:
        number = -number
    return number


def parse_date_value(value: str, with_time: bool = False) -> date | datetime | None:
    text = value.strip()
    if not text or text in {"-", "--"}:
        return None
    formats = (
        ["%d/%m/%Y %H:%M:%S", "%d/%m/%Y %H:%M", "%Y-%m-%dT%H:%M:%S", "%Y-%m-%d %H:%M:%S"]
        if with_time
        else ["%d/%m/%Y", "%Y-%m-%d", "%d-%m-%Y"]
    )
    for fmt in formats:
        try:
            parsed = datetime.strptime(text, fmt)
            return parsed if with_time else parsed.date()
        except ValueError:
            continue
    raise ValueError(f"Data invalida: {value!r}")


def safe_excel_text(value: str) -> str:
    """Neutraliza formula injection em textos extraidos de documentos externos."""
    if value.lstrip().startswith(("=", "+", "-", "@")):
        return "'" + value
    return value


def coerce_value(raw: str, data_type: DataType):
    value = raw.strip()
    if not value:
        return None
    if data_type == "text":
        return safe_excel_text(value)
    if data_type == "integer":
        number = parse_decimal(value, integer=True)
        if number is not None and number != number.to_integral_value():
            raise ValueError("Quantidade inteira contém casas decimais.")
        return int(number) if number is not None else None
    if data_type in {"decimal", "currency"}:
        number = parse_decimal(value)
        return float(number) if number is not None else None
    if data_type == "percentage":
        number = parse_decimal(value)
        if number is None:
            return None
        if "%" in value or abs(number) > 1:
            number /= Decimal("100")
        return float(number)
    if data_type == "date":
        return parse_date_value(value, with_time=False)
    if data_type == "datetime":
        return parse_date_value(value, with_time=True)
    if data_type == "boolean":
        normalized = strip_accents(value).casefold()
        if normalized in {"sim", "true", "verdadeiro", "1"}:
            return True
        if normalized in {"nao", "false", "falso", "0"}:
            return False
        raise ValueError(f"Booleano invalido: {raw!r}")
    return safe_excel_text(value)


def number_format_for(column: ColumnDefinition) -> str:
    defaults = {
        "text": "@",
        "integer": "#,##0",
        "decimal": "#,##0.00",
        "currency": "#,##0.00",
        "percentage": "0.00%",
        "date": "dd/mm/yyyy",
        "datetime": "dd/mm/yyyy hh:mm",
        "boolean": "General",
    }
    # O formato recebido do modelo e tratado como dado nao confiavel. So aceitamos
    # uma lista curta de formatos conhecidos para evitar formulas ou strings
    # arbitrarias dentro do arquivo final.
    allowed_formats = {
        "@",
        "General",
        "#,##0",
        "#,##0.00",
        "#,##0.000",
        "0.00%",
        "0.0%",
        "dd/mm/yyyy",
        "dd/mm/yyyy hh:mm",
    }
    suggested = column.number_format.strip()
    return suggested if suggested in allowed_formats else defaults[column.data_type]


def source_line(extraction: DocumentExtraction, table: ExtractedTable) -> str:
    base = f"Fonte: {extraction.source_file} - conteudo visivel no PDF."
    details: list[str] = []
    if table.source_pages:
        pages = ", ".join(str(page) for page in table.source_pages)
        details.append(f"Pagina(s): {pages}.")
    if table.source_note.strip():
        details.append(table.source_note.strip().rstrip(".") + ".")
    details.extend(item.strip().rstrip(".") + "." for item in table.limitations if item.strip())
    return " ".join([base, *details])


def apply_full_row_fill(ws, row: int, max_column: int, color: str) -> None:
    fill = PatternFill("solid", fgColor=color)
    for column in range(1, max_column + 1):
        ws.cell(row=row, column=column).fill = fill


def apply_row_kind_style(ws, row: int, kind: RowKind, max_column: int) -> None:
    if kind == "total":
        fill = PatternFill("solid", fgColor=TOTAL_GRAY)
        for column in range(1, max_column + 1):
            cell = ws.cell(row=row, column=column)
            cell.fill = fill
            cell.font = Font(name="Carlito", size=10, bold=True, color="111827")
    elif kind == "subtotal":
        top = Side(style="thin", color=BORDER_GRAY)
        for column in range(1, max_column + 1):
            cell = ws.cell(row=row, column=column)
            cell.font = Font(name="Carlito", size=10, bold=True, color=TEXT_DARK)
            cell.border = Border(top=top)
    elif kind == "note":
        for column in range(1, max_column + 1):
            ws.cell(row=row, column=column).font = Font(
                name="Carlito", size=10, italic=True, color=TEXT_MUTED
            )


def set_column_widths(ws, table: ExtractedTable) -> None:
    for col_index, column in enumerate(table.columns, start=1):
        lengths = [len(column.name)]
        for row in table.rows:
            if col_index - 1 < len(row.values):
                lengths.append(len(str(row.values[col_index - 1])))
        visible_max = max(lengths, default=10)
        if column.data_type == "text":
            width = min(max(12, visible_max + 2), 70)
        elif column.data_type in {"date", "datetime"}:
            width = min(max(14, visible_max + 2), 22)
        else:
            width = min(max(12, visible_max + 2), 28)
        ws.column_dimensions[get_column_letter(col_index)].width = width


def style_cell_by_type(cell: Cell, column: ColumnDefinition) -> None:
    cell.number_format = number_format_for(column)
    cell.font = Font(name="Carlito", size=10, color=TEXT_DARK)
    cell.alignment = Alignment(
        horizontal="left" if column.data_type == "text" else "right",
        vertical="center",
        wrap_text=column.data_type == "text",
    )


def build_workbook(
    extraction: DocumentExtraction,
    output_path: Path,
    freeze: bool = True,
) -> None:
    workbook = Workbook()
    workbook.remove(workbook.active)
    workbook.properties.creator = "Fundacred - Extrator de Recebiveis"
    workbook.properties.title = f"Tabelas extraidas de {extraction.source_file}"
    workbook.properties.description = (
        "Uma aba por tabela visivel no PDF. Valores preservados conforme a fonte."
    )

    used_sheet_names: set[str] = set()
    used_table_names: set[str] = set()

    if not extraction.tables:
        ws = workbook.create_sheet("Sem Tabelas")
        ws.sheet_view.showGridLines = False
        ws["A1"] = "Nenhuma tabela encontrada"
        ws["A1"].fill = PatternFill("solid", fgColor=FUNDACRED_BLUE)
        ws["A1"].font = Font(name="Carlito", size=16, bold=True, color=WHITE)
        ws["A2"] = f"Fonte: {extraction.source_file}"
        ws["A4"] = "Revise o PDF e os avisos registrados no JSON de auditoria."
        ws.column_dimensions["A"].width = 70
    else:
        for index, extracted_table in enumerate(extraction.tables, start=1):
            sheet_name = unique_sheet_name(
                extracted_table.sheet_name or extracted_table.title,
                used_sheet_names,
            )
            ws = workbook.create_sheet(sheet_name)
            ws.sheet_view.showGridLines = False
            max_column = len(extracted_table.columns)
            last_column = get_column_letter(max_column)

            # Faixa de titulo.
            for column_index in range(1, max_column + 1):
                cell = ws.cell(row=1, column=column_index)
                cell.fill = PatternFill("solid", fgColor=FUNDACRED_BLUE)
            ws.merge_cells(f"A1:{last_column}1")
            ws["A1"] = safe_excel_text(extracted_table.title)
            ws["A1"].font = Font(name="Carlito", size=16, bold=True, color=WHITE)
            ws["A1"].alignment = Alignment(horizontal="left", vertical="center")
            ws.row_dimensions[1].height = 30

            # Faixa de fonte e limitacoes.
            for column_index in range(1, max_column + 1):
                ws.cell(row=2, column=column_index).fill = PatternFill(
                    "solid", fgColor=SOURCE_BLUE
                )
            ws.merge_cells(f"A2:{last_column}2")
            note = source_line(extraction, extracted_table)
            ws["A2"] = note
            ws["A2"].font = Font(
                name="Carlito", size=9, italic=True, color=TEXT_MUTED
            )
            ws["A2"].alignment = Alignment(vertical="center", wrap_text=True)
            estimated_lines = max(1, math.ceil(len(note) / max(45, max_column * 18)))
            ws.row_dimensions[2].height = min(72, 16 + estimated_lines * 9)

            # Cabecalho da tabela na linha 4.
            header_bottom = Side(style="thin", color=BORDER_GRAY)
            used_headers: set[str] = set()
            for col_index, column in enumerate(extracted_table.columns, start=1):
                header = safe_excel_text(column.name.strip() or f"Coluna {col_index}")
                base_header = header
                suffix = 2
                while header.casefold() in used_headers:
                    header = f"{base_header} ({suffix})"
                    suffix += 1
                used_headers.add(header.casefold())
                cell = ws.cell(row=4, column=col_index, value=header)
                cell.fill = PatternFill("solid", fgColor=HEADER_GRAY)
                cell.font = Font(name="Carlito", size=11, bold=True, color=TEXT_DARK)
                cell.alignment = Alignment(
                    horizontal="center", vertical="center", wrap_text=True
                )
                cell.border = Border(bottom=header_bottom)
            ws.row_dimensions[4].height = 24

            # Dados tipados.
            for row_offset, extracted_row in enumerate(extracted_table.rows, start=5):
                for col_index, (raw, column) in enumerate(
                    zip(extracted_row.values, extracted_table.columns), start=1
                ):
                    cell = ws.cell(row=row_offset, column=col_index)
                    try:
                        cell.value = coerce_value(raw, column.data_type)
                    except ValueError:
                        # Preserva o valor legivel e sinaliza o erro na propria celula.
                        cell.value = safe_excel_text(raw)
                        cell.number_format = "@"
                        cell.font = Font(
                            name="Carlito", size=10, italic=True, color="C23A3A"
                        )
                    else:
                        style_cell_by_type(cell, column)
                ws.row_dimensions[row_offset].height = 18
                apply_row_kind_style(
                    ws, row_offset, extracted_row.row_kind, max_column
                )

            last_row = max(4, 4 + len(extracted_table.rows))
            if extracted_table.columns:
                table_name = unique_table_name(sheet_name, index, used_table_names)
                excel_table = Table(
                    displayName=table_name,
                    ref=f"A4:{last_column}{last_row}",
                )
                excel_table.tableStyleInfo = TableStyleInfo(
                    name="TableStyleMedium2",
                    showFirstColumn=False,
                    showLastColumn=False,
                    showRowStripes=True,
                    showColumnStripes=False,
                )
                ws.add_table(excel_table)

            set_column_widths(ws, extracted_table)
            if freeze:
                ws.freeze_panes = "A5"
            ws.auto_filter.ref = f"A4:{last_column}{last_row}"
            ws.print_title_rows = "1:4"
            ws.sheet_properties.pageSetUpPr.fitToPage = True
            ws.page_setup.fitToWidth = 1
            ws.page_setup.fitToHeight = 0
            ws.page_setup.orientation = "landscape" if max_column > 5 else "portrait"
            ws.page_margins.left = 0.25
            ws.page_margins.right = 0.25
            ws.page_margins.top = 0.35
            ws.page_margins.bottom = 0.35

    output_path.parent.mkdir(parents=True, exist_ok=True)
    workbook.save(output_path)


def verify_workbook(output_path: Path, expected_tables: int) -> None:
    workbook = load_workbook(output_path, read_only=False, data_only=False)
    expected_sheets = expected_tables if expected_tables else 1
    if len(workbook.sheetnames) != expected_sheets:
        raise RuntimeError(
            f"XLSX invalido: esperadas {expected_sheets} abas, "
            f"encontradas {len(workbook.sheetnames)}."
        )
    for ws in workbook.worksheets:
        if ws.max_row < 1 or ws.max_column < 1:
            raise RuntimeError(f"Aba vazia encontrada: {ws.title}")
        if ws.title != "Sem Tabelas":
            if ws["A1"].fill.fgColor.rgb != FUNDACRED_BLUE:
                raise RuntimeError(f"Titulo sem estilo Fundacred na aba {ws.title}.")
            if ws.max_row >= 5 and not ws.tables:
                raise RuntimeError(f"Tabela estruturada ausente na aba {ws.title}.")
    workbook.close()


def save_json(extraction: DocumentExtraction, json_path: Path) -> None:
    json_path.parent.mkdir(parents=True, exist_ok=True)
    json_path.write_text(
        json.dumps(extraction.model_dump(), ensure_ascii=False, indent=2),
        encoding="utf-8",
    )


def main() -> int:
    args = parse_args()
    output_path = args.output or default_output_path(args.pdf, args.json_input)

    if args.pdf:
        page_count = validate_pdf(args.pdf, args.max_pages, args.max_mb)
        extraction = extract_with_gemini(
            pdf_path=args.pdf,
            model=args.model,
            expected_pages=page_count,
            attempts=args.attempts,
        )
    else:
        extraction = load_extraction_json(args.json_input)

    if args.save_json:
        save_json(extraction, args.save_json)

    build_workbook(extraction, output_path, freeze=not args.no_freeze)
    verify_workbook(output_path, len(extraction.tables))

    result = {
        "status": "CONCLUIDO_COM_AVISOS" if extraction.warnings else "CONCLUIDO",
        "arquivo": str(output_path.resolve()),
        "fonte": extraction.source_file,
        "paginas": extraction.page_count,
        "tabelas": len(extraction.tables),
        "avisos": extraction.warnings,
    }
    print(json.dumps(result, ensure_ascii=False, indent=2))
    return 0


if __name__ == "__main__":
    try:
        raise SystemExit(main())
    except (ValidationError, ValueError, FileNotFoundError, RuntimeError) as exc:
        print(
            json.dumps(
                {"status": "ERRO", "erro": str(exc)},
                ensure_ascii=False,
                indent=2,
            ),
            file=sys.stderr,
        )
        raise SystemExit(1)
