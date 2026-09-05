from __future__ import annotations

import sys
import tempfile
import unittest
from pathlib import Path

from openpyxl import load_workbook

PROJECT_ROOT = Path(__file__).resolve().parents[1]
sys.path.insert(0, str(PROJECT_ROOT))

from extrair_recebiveis_gemini import (  # noqa: E402
    FUNDACRED_BLUE,
    build_workbook,
    coerce_value,
    load_extraction_json,
    verify_workbook,
)


class ConversionTests(unittest.TestCase):
    def test_brazilian_values(self) -> None:
        self.assertEqual(coerce_value("R$ 1.234,56", "currency"), 1234.56)
        self.assertAlmostEqual(coerce_value("4,34%", "percentage"), 0.0434)
        self.assertEqual(coerce_value("1.234", "integer"), 1234)
        self.assertEqual(coerce_value("00123", "text"), "00123")

    def test_formula_injection_is_neutralized(self) -> None:
        self.assertEqual(coerce_value("=2+2", "text"), "'=2+2")


class WorkbookTests(unittest.TestCase):
    def test_example_workbook(self) -> None:
        extraction = load_extraction_json(PROJECT_ROOT / "exemplo_extracao.json")
        with tempfile.TemporaryDirectory() as temp_dir:
            output = Path(temp_dir) / "resultado.xlsx"
            build_workbook(extraction, output)
            verify_workbook(output, len(extraction.tables))

            workbook = load_workbook(output, data_only=False)
            self.assertEqual(
                workbook.sheetnames,
                ["Resumo Processos", "Inadimplencia", "Recebiveis"],
            )
            for worksheet in workbook.worksheets:
                self.assertEqual(worksheet["A1"].fill.fgColor.rgb, FUNDACRED_BLUE)
                self.assertEqual(worksheet.freeze_panes, "A5")
                self.assertTrue(worksheet.tables)
            self.assertIsInstance(workbook["Resumo Processos"]["B5"].value, int)
            self.assertIsInstance(workbook["Inadimplencia"]["B5"].value, float)
            workbook.close()


if __name__ == "__main__":
    unittest.main()
