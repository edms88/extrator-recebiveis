import tempfile
import unittest
from pathlib import Path
from unittest.mock import patch, MagicMock
from datetime import date

from openpyxl import load_workbook
from pypdf import PdfWriter
from extrair_recebiveis_gemini import (
    coerce_value, load_extraction_json, build_workbook, validate_pdf,
    UserFacingError,
)

ROOT = Path(__file__).resolve().parents[1]


class SecurityTests(unittest.TestCase):
    def test_all_text_surfaces_and_empty_table(self):
        extraction = load_extraction_json(ROOT / 'exemplo_extracao.json')
        for table in extraction.tables:
            table.title = '=1+1'
            table.sheet_name = 'Mesmo/nome'
            for column in table.columns:
                column.name = '=1+1'
        extraction.tables[-1].rows = []
        with tempfile.TemporaryDirectory() as directory:
            path = Path(directory) / 'test.xlsx'
            build_workbook(extraction, path)
            wb = load_workbook(path)
            self.assertEqual(len(set(wb.sheetnames)), 3)
            for ws in wb:
                self.assertTrue(ws.tables)
                headers = [c.value for c in ws[4]]
                self.assertEqual(len(headers), len(set(headers)))
                self.assertFalse(any(c.data_type == 'f' for row in ws for c in row))
            wb.close()

    def test_types(self):
        self.assertEqual(coerce_value('31/08/2026', 'date'), date(2026, 8, 31))
        self.assertEqual(coerce_value('-12,30', 'currency'), -12.3)
        for prefix in '=+-@':
            self.assertTrue(coerce_value(prefix + 'abc', 'text').startswith("'"))

    def test_pdf_validation(self):
        with tempfile.TemporaryDirectory() as directory:
            path = Path(directory) / 'test.pdf'
            path.write_bytes(b'not pdf')
            with self.assertRaises(ValueError):
                validate_pdf(path, 10, 10)
            writer = PdfWriter()
            for _ in range(11):
                writer.add_blank_page(width=100, height=100)
            writer.write(path)
            with self.assertRaises(ValueError):
                validate_pdf(path, 10, 10)
            self.assertEqual(validate_pdf(path, 11, 10), 11)



if __name__ == '__main__':
    unittest.main()
